"""
Модуль экспорта отчётов о расчёте отделочных материалов.

Принимает Summary из calculator.py и сохраняет его в .docx и/или .xlsx.
"""

from __future__ import annotations

from abc import ABC, abstractmethod
from datetime import datetime
from pathlib import Path

# python-docx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

# openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .calculator import CalculationResult, Summary


# Заголовки таблицы
TABLE_HEADERS = (
    "Материал",
    "Площадь единицы, м²",
    "Единица",
    "Площадь, м²",
    "Запас",
    "Кол-во",
    "Цена, ₽",
    "Стоимость, ₽",
)


class BaseReportWriter(ABC):
    """
    Абстрактный базовый класс для экспортёров отчёта.

    Наследники (DocxReportWriter, XlsxReportWriter) реализуют
    _write_report — работу с конкретным форматом файла.
    Общая логика (создание папки, имя файла, дата) — здесь.
    """

    # Расширение файла — определяет наследник.
    EXTENSION: str = ""

    def __init__(self, output_dir: str | Path = "reports"):
        self._output_dir = None
        self.output_dir = output_dir

    @property
    def output_dir(self) -> Path:
        return self._output_dir

    @output_dir.setter
    def output_dir(self, value: str | Path):
        if isinstance(value, (str, Path)):
            self._output_dir = Path(value)
        else:
            raise TypeError("output_dir должен быть строкой или Path")

    # ---------- публичный API ----------

    def save(
        self,
        summary: Summary,
        filename: str | None = None,
    ) -> Path:
        """
        Сохранить отчёт и вернуть путь к файлу.
        """
        if not isinstance(summary, Summary):
            raise TypeError("Ожидался объект Summary")

        # Создаём папку, если её ещё нет (parents=True — на случай вложенных путей).
        self._output_dir.mkdir(parents=True, exist_ok=True)

        if filename is None:
            timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
            filename = f"report_{timestamp}.{self.EXTENSION}"

        filepath = self._output_dir / filename

        self._write_report(summary, filepath)
        return filepath

    @abstractmethod
    def _write_report(self, summary: Summary, filepath: Path) -> None:
        """Реализуется в наследнике: записать файл конкретного формата."""
        ...

    def __repr__(self) -> str:
        return f"{type(self).__name__}(output_dir={self.output_dir!r})"


# =====================================================================
#                            DOCX
# =====================================================================

class DocxReportWriter(BaseReportWriter):
    """Сохраняет отчёт в .docx через python-docx."""

    EXTENSION = "docx"

    def _write_report(self, summary: Summary, filepath: Path) -> None:
        doc = Document()

        # --- заголовок ---
        title = doc.add_heading("Отчёт о расчёте отделочных материалов", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # --- метаданные ---
        date_par = doc.add_paragraph()
        date_par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        date_run = date_par.add_run(
            f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        )
        date_run.font.size = Pt(10)
        date_run.italic = True

        # --- таблица ---
        table = doc.add_table(rows=1, cols=len(TABLE_HEADERS))
        # Встроенный стиль с границами и заливкой шапки
        table.style = "Light Grid Accent 1"

        header_cells = table.rows[0].cells
        for i, header in enumerate(TABLE_HEADERS):
            cell = header_cells[i]
            cell.text = header
            # Жирный шрифт в шапке
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

        # --- строки данных ---
        for result in summary.results:
            row_cells = table.add_row().cells
            for i, value in enumerate(result.to_row()):
                row_cells[i].text = str(value)

        # --- итог ---
        doc.add_paragraph()
        total_par = doc.add_paragraph()
        total_run = total_par.add_run(f"ИТОГО: {summary.total:.2f} ₽")
        total_run.bold = True
        total_run.font.size = Pt(14)

        doc.save(str(filepath))

    def __str__(self) -> str:
        return f"Экспортёр .docx (папка: {self.output_dir})"


# =====================================================================
#                            XLSX
# =====================================================================

class XlsxReportWriter(BaseReportWriter):
    """Сохраняет отчёт в .xlsx через openpyxl."""

    EXTENSION = "xlsx"

    # Стили вынесены в атрибуты класса — создаются один раз, не на каждый файл.
    HEADER_FILL = PatternFill("solid", fgColor="4472C4")
    HEADER_FONT = Font(bold=True, color="FFFFFF")
    TOTAL_FONT = Font(bold=True, size=12)
    CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
    THIN = Side(style="thin", color="999999")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    def _write_report(self, summary: Summary, filepath: Path) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Расчёт"

        # --- заголовок отчёта ---
        ws["A1"] = "Отчёт о расчёте отделочных материалов"
        ws["A1"].font = Font(bold=True, size=14)

        ws["A2"] = f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        ws["A2"].font = Font(italic=True, size=10)

        # --- шапка таблицы (строка 4) ---
        header_row = 4
        for col_idx, header in enumerate(TABLE_HEADERS, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=header)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.alignment = self.CENTER
            cell.border = self.BORDER

        # --- строки данных ---
        for offset, result in enumerate(summary.results, start=1):
            row = header_row + offset
            for col_idx, value in enumerate(result.to_row(), start=1):
                cell = ws.cell(row=row, column=col_idx, value=value)
                cell.border = self.BORDER

        # --- строка "ИТОГО" ---
        total_row = header_row + len(summary.results) + 1
        label_cell = ws.cell(row=total_row, column=1, value="ИТОГО")
        label_cell.font = self.TOTAL_FONT

        total_cell = ws.cell(
            row=total_row,
            column=len(TABLE_HEADERS),
            value=round(summary.total, 2),
        )
        total_cell.font = self.TOTAL_FONT
        total_cell.number_format = '#,##0.00 "₽"'
        
        for col_idx in range(1, len(TABLE_HEADERS) + 1):
            column_letter = ws.cell(row=1, column=col_idx).column_letter
            max_len = 0
            for row in range(header_row, total_row + 1):
                value = ws.cell(row=row, column=col_idx).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            ws.column_dimensions[column_letter].width = max_len + 2

        wb.save(str(filepath))

    def __str__(self) -> str:
        return f"Экспортёр .xlsx (папка: {self.output_dir})"


# =====================================================================
#                     ФУНКЦИЯ-ОБЁРТКА
# =====================================================================

def save_all(
    summary: Summary,
    output_dir: str | Path = "reports",
    docx: bool = True,
    xlsx: bool = True,
) -> list[Path]:
    """
    Сохранить отчёт во все запрошенные форматы за один вызов.

    Возвращает список путей к созданным файлам.
    """
    saved: list[Path] = []
    if docx:
        saved.append(DocxReportWriter(output_dir).save(summary))
    if xlsx:
        saved.append(XlsxReportWriter(output_dir).save(summary))
    return saved